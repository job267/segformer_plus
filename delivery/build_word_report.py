"""Merge the completed experiment records into a single editable course report."""

import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports/SegFormer_B0_课程设计与实验综合报告.docx"
FIGURES = ROOT / "reports/figures"
NAMES = ["Sky", "Building", "Pole", "Road", "Sidewalk", "Tree", "SignSymbol", "Fence", "Car", "Pedestrian", "Bicyclist"]
CN = ["天空", "建筑", "杆状物", "道路", "人行道", "树木", "标志牌", "围栏", "车辆", "行人", "骑行者"]
RUNS = ["E0_baseline", "E1_augmentation", "E2_dice", "E3_combined", "E4_resolution", "E5_weighted",
        "multiseed/E4_seed2024", "multiseed/E5_seed2024", "multiseed/E4_seed3407", "multiseed/E5_seed3407"]
SOURCE_DOCS = ["reports/final_report.md", "reports/baseline_analysis.md", "reports/ablation.md",
               "reports/resolution_ablation.md", "reports/weighted_ablation.md", "reports/multiseed.md",
               "reports/inference_benchmark.md", "reports/stage0_stage1_status.md", "docs/reproduction.md"]


def read_json(relative):
    return json.loads((ROOT / relative).read_text(encoding="utf-8"))


def digest(path):
    value = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024**2), b""):
            value.update(block)
    return value.hexdigest()


def md_table(headers, rows):
    return "\n".join(["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers),
                      *["| " + " | ".join(map(str, row)) + " |" for row in rows]])


def verified_inputs():
    """Independent archive check with no training dependency or dataset inference."""
    preflight = read_json("results/multiseed/preflight.json")
    for name, expected in preflight["source_sha256"].items():
        if digest(ROOT / name) != expected:
            raise ValueError(f"Frozen source changed: {name}")
    records = []
    for name in RUNS:
        directory = ROOT / "results" / name
        metrics = json.loads((directory / "metrics.json").read_text(encoding="utf-8"))
        history = list(csv.DictReader((directory / "train_log.csv").open(encoding="utf-8")))
        if [int(row["epoch"]) for row in history] != list(range(1, 51)):
            raise ValueError(f"Incomplete history: {name}")
        best = max(history, key=lambda row: float(row["val_miou"]))
        if int(best["epoch"]) != metrics["best_val_epoch"] or abs(float(best["val_miou"])-metrics["best_val_miou"]) > 1e-12:
            raise ValueError(f"Validation selection changed: {name}")
        matrix = metrics["confusion_matrix"]
        totals = [sum(row) for row in matrix]
        ious = [matrix[i][i] / (totals[i] + sum(row[i] for row in matrix) - matrix[i][i]) for i in range(11)]
        mean = sum(ious)/11
        accuracy = sum(matrix[i][i] for i in range(11))/sum(totals)
        if abs(mean-metrics["test_miou"]) > 1e-12 or abs(accuracy-metrics["pixel_accuracy"]) > 1e-12:
            raise ValueError(f"Test confusion differs: {name}")
        if any(abs(a-b) > 1e-12 for a,b in zip(ious, metrics["per_class_iou"])):
            raise ValueError(f"Class IoU differs: {name}")
        if digest(directory / "best_model/model.safetensors") != metrics["checkpoint_sha256"]:
            raise ValueError(f"Checkpoint differs: {name}")
        journal = json.loads((directory / "test_evaluation.json").read_text(encoding="utf-8"))
        if sum(bool(attempt.get("completed_utc")) for attempt in journal["attempts"]) != 1:
            raise ValueError(f"Test journal differs: {name}")
        records.append({"name":name, "metrics":metrics, "history":history})
    return records


def merged_markdown(records):
    source = (ROOT / "reports/final_report.md").read_text(encoding="utf-8")
    parts = re.split(r"(?m)^## ", source)
    sections = {part.splitlines()[0]: "\n".join(part.splitlines()[1:]).strip() for part in parts[1:]}
    get = lambda prefix: next(value for key,value in sections.items() if key.startswith(prefix))
    audit = read_json("results/stage1/data_audit.json")
    overfit = read_json("results/stage1/overfit/overfit_check.json")
    weights = read_json("results/stage6/class_weights.json")["weights"]
    metrics = [record["metrics"] for record in records]
    pct = lambda value: f"{100*value:.2f}%"
    output = "# 基于 SegFormer B0 的 CamVid 街景语义分割\n\n课程设计与实验综合报告\n\n图像处理与分析\n\n2026年10月19日\n\n## 摘要\n\n" + get("摘要")
    output += "\n\n[[TOC]]\n\n## 1 任务目标与研究范围\n\n" + get("1.")
    basis=get("2.").replace("每类 `IoU_c = TP_c / (TP_c + FP_c + FN_c)`", "每类 IoU 按下式计算")
    output += "\n\n## 2 模型方法与评价指标\n\n" + basis + "\n\n[[EQUATION:IOU]]"
    output += "\n\n### 损失函数与类别平衡\n\n设有效像素集合为 V，类别 c 的预测概率与 one hot 标签分别为 p 和 g。忽略区域先从概率与标签中屏蔽。交叉熵对有效像素归约，E5 则按有效像素类别权重之和归一化；所有权重为 1 时还原普通 CE。\n\n[[EQUATION:CE]]\n\nDice 在整个批次和空间维度累计，仅对该批 GT 实际出现的类别集合 C 取平均，平滑项 ε=10⁻⁶。GT 未出现的类不参与 Dice 均值，但 CE 仍监督其预测。\n\n[[EQUATION:DICE]]\n\nE2 至 E5 的组合损失采用同一系数。\n\n[[EQUATION:LOSS]]\n\n类别权重先用未增强训练标签的有效像素频率取逆平方根，再除以类别均值并截断。频率统计、损失中的有效像素归约和用于 mIoU 的混淆矩阵各自保持明确口径。\n"
    output += "\n## 3 数据环境与基础验收\n\n" + get("3.")
    output += "\n\n### 数据划分与类别编码\n\n" + md_table(["划分", "图像数量", "原始尺寸", "原始 ignore 比例"],
        [[label, audit["splits"][split]["audited_samples"], "960×720", f"{100*audit['splits'][split]['ignored_pixels']/audit['splits'][split]['total_pixels']:.4f}%"] for split,label in (("train","训练"),("val","验证"),("test","测试"))])
    output += "\n\n表中 ignore 比例属于原始分辨率，不是缩放后评价网格的比例。完整文件与清单哈希、跨划分重用检查保留在数据审计记录。\n\n" + md_table(["ID", "类别", "中文含义", "E5 CE 权重"], [[i,name,CN[i],f"{weights[i]:.4f}"] for i,name in enumerate(NAMES)])
    output += "\n\n255 为无效标注，不属于第十二个预测类别，也不计入权重和指标。权重最小截断值为 0.5，最大允许值为 3.0；上表为实际使用值。\n\n### 基础构建与学习能力验证\n\n正式训练前的关键验收如下，诊断结果与正式 Val/Test 成绩严格区分。\n\n" + md_table(["检查", "实际证据", "用途"], [
        ["完整数据审计", "367 / 101 / 233 全部读取且错误数 0", "核对配对 标签与划分"],
        ["预训练载入", "编码器严格复制 仅丢弃原分类头两项", "避免静默随机初始化"],
        ["真实 CUDA 批次", "4 张样本 CE 2.403232 参数成功更新", "检查前向 反向与优化器"],
        ["小样本学习诊断", "20 张训练图 20 轮 学习率 0.001", "验证流程具有学习能力"],
        ["交付核心测试", "77 项通过", "覆盖指标 损失 恢复与交付完整性"]])
    output += f"\n\n小样本诊断的子集 CE 从 {overfit['initial']['loss']:.6f} 降至 {overfit['final']['loss']:.6f}，相对下降 {100*overfit['relative_loss_decrease']:.2f}%；同一训练子集 mIoU 从 {pct(overfit['initial']['miou'])} 上升至 {pct(overfit['final']['miou'])}。该诊断满足预先约定的学习检查，但末轮仍有未学好的小类别，不等同于充分过拟合或泛化验证。其权重没有用于正式训练初始化。\n"
    output += "\n## 4 实验设计与运行控制\n\n" + get("4.")
    output += "\n\n### 单种子实验矩阵\n\n" + md_table(["组别", "训练网格", "成对增强", "损失", "主要对照"],[
        ["E0","384×384","关闭","CE","初始基线"],["E1","384×384","开启","CE","E0 增强"],
        ["E2","384×384","关闭","CE + 0.5 Dice","E0 损失"],["E3","384×384","开启","CE + 0.5 Dice","合并两项"],
        ["E4","512×512","开启","CE + 0.5 Dice","E3 分辨率"],["E5","512×512","开启","加权 CE + 0.5 Dice","E4 类别监督"]])
    output += "\n\n### 选模与复核规则\n\nE0 至 E3 构成原定四组消融；E4 和 E5 为后续独立扩展。各组最佳轮次都由验证 mIoU 决定，验证结束后才评价所选模型。E4/E5 的配对复核预先固定 42、2024、3407；同种子初始完整模型哈希相同。四次新增运行只改变种子与输出身份，保留原始实现、训练频率权重和预算。\n\n由于后续扩展是在已查看历史测试结果后提出，整个项目不具备完全未见测试集的盲测条件。该限制影响最终结论的外推范围，即使单次训练内部始终没有用 Test 选轮次，也不能省略。\n"
    result_section = get("5.")
    # Tables are rebuilt from all six actual metrics; the course draft showed only four class columns.
    result_section = re.sub(r"\| 类别 \| E0 Test IoU.*?(?=\n\nE5 相对)", md_table(["类别","E0","E1","E2","E3","E4","E5"], [[name,*[pct(item["per_class_iou"][i]) for item in metrics[:6]]] for i,name in enumerate(NAMES)]), result_section, flags=re.S)
    result_section = result_section.replace("![E0 逐轮", "### 收敛曲线与混淆分析\n\n![E0 逐轮", 1)
    output += "\n## 5 单种子实验结果与类别分析\n\n" + result_section
    e0 = records[0]
    first,last=e0["history"][0],e0["history"][-1]
    total_time=sum(float(row["epoch_time_sec"]) for row in e0["history"])/60
    output += f"\n\n### 基线的收敛与泛化差距\n\nE0 的训练 CE 从首轮 {float(first['train_loss']):.6f} 降至第 50 轮 {float(last['train_loss']):.6f}，验证 CE 从 {float(first['val_loss']):.6f} 降至 {float(last['val_loss']):.6f}。末轮验证 mIoU 为 {pct(float(last['val_miou']))}，距最佳轮次只有 {100*(metrics[0]['best_val_miou']-float(last['val_miou'])):.2f} 个百分点，后期趋于平台，没有持续性严重退化的证据。训练与验证逐轮累计 {total_time:.2f} 分钟，不含启动与归档开销。\n\n最佳验证与测试 mIoU 相差 {100*(metrics[0]['best_val_miou']-metrics[0]['test_miou']):.2f} 个百分点；场景与类别组成差异也可能参与其中，不能仅凭差距把原因归结为过拟合。以真实 Pole 像素为分母，E0 中 {100*metrics[0]['confusion_matrix'][2][1]/sum(metrics[0]['confusion_matrix'][2]):.2f}% 被预测为 Building。该比例来自混淆矩阵的 GT 行，与 Pole IoU 的定义不同。\n"
    output += "\n## 6 多随机种子复核与最终模型\n\n" + get("6.")
    output += "\n\n### 六次配对运行的实际结果\n\n" + md_table(["方案", "种子", "最佳轮次", "Val mIoU", "Test mIoU", "像素准确率"], [[m["experiment_id"].split("_")[0],m["seed"],m["best_val_epoch"],pct(m["best_val_miou"]),pct(m["test_miou"]),pct(m["pixel_accuracy"])] for m in [metrics[i] for i in (4,5,6,7,8,9)]])
    def stats(values):
        avg=sum(values)/len(values)
        return avg,(sum((x-avg)**2 for x in values)/(len(values)-1))**0.5
    perclass=[]
    for i,name in enumerate(NAMES):
        e4=stats([metrics[j]["per_class_iou"][i] for j in (4,6,8)])
        e5=stats([metrics[j]["per_class_iou"][i] for j in (5,7,9)])
        perclass.append([name,f"{pct(e4[0])} ± {100*e4[1]:.2f}",f"{pct(e5[0])} ± {100*e5[1]:.2f}",f"{100*(e5[0]-e4[0]):+.2f}"])
    output += "\n\n### 三种子的逐类测试表现\n\n" + md_table(["类别", "E4 均值±std", "E5 均值±std", "E5−E4 均值 pp"],perclass)
    output += "\n\n表中均值以百分比表示，std 和差值单位为百分点。逐类统计与整体 mIoU 共同描述类别取舍；较高的像素准确率不意味着小类别一定更好，也不将任意类别的一次改善视为普遍机制。\n"
    output += "\n## 7 定性结果与失败案例\n\n两个样本来自预先固定的六个测试清单位置，均展示原图、GT、E0 与 E5。图片为原始 960×720 网格；量化指标仍在共同 384×384 标签网格计算。显示图不能当作原分辨率成绩，也不能代表全部测试图。\n\n### 街景样本一\n\n[[CASE:0001TP_008550]]\n\nE5 的左侧杆状结构比 E0 更连续，右侧保留了一个较小的人物区域，但建筑和植被边界仍有碎片。这些是同图观察，没有为单张样例另算分数。\n\n### 街景样本二\n\n[[CASE:Seq05VD_f02310]]\n\nGT 中细枝之间的间隙在预测中被较大植被区域填充，部分杆与标志物轮廓增粗，右侧远处人物未被完整分开。道路主体得到识别的同时，复杂细节仍然失败。保留这类结果有助于避免仅展示容易样本。\n"
    output += "\n## 8 推理系统与效率评测\n\n" + get("8.")
    output += "\n\n## 9 工程复现与交付验收\n\n" + get("9.")
    output += "\n\n### 离线迁移验证\n\n交付版本在独立新目录解压后，346 个清单文件的大小与 SHA256 全部相符，十次完整训练归档和最终验证均值选模通过核验。在没有完整 CamVid 和预训练缓存的条件下，强制 Hugging Face 离线，使用 CPU 对随附的一张验证示例完成 FP32 预测，输出为 960×720，类别编号均在 0 至 10。另一个不含历史 results 的纯源码目录通过全部 77 项测试。\n\n上述检查使用本机已有固定依赖，没有在另一台新机器重新安装依赖，也没有重新训练十次实验；它验证的是文件完整性、目录迁移和离线推理能力。完整数据集与最后训练状态没有进入提交 ZIP，因而该 ZIP 不是原训练现场的全量备份。\n\n### 实施中的问题与处理\n\n数据获取曾发生连接中断，下载器通过持久连接、重试和已有文件复用完成获取，没有关闭证书校验或替换来源。并发下载的目标路径改为在线程启动前统一解析，继续保留越界拒绝检查。为避开当前 PyTorch 版本融合 CUDA NLL 空间归约的确定性警告，CE 采用逐像素计算和显式有效像素归约，并核对手算结果。Windows 无符号链接权限时缓存使用普通文件，不需要修改系统权限。\n\n这些处理服务于可核查的输入、计算和恢复过程，不能被当作额外精度提升来源。\n"
    output += "\n## 10 结论与课程实践认识\n\n" + get("10.")
    output += "\n\n### 课程实践认识\n\n本项目把图像缩放、像素编码、成对变换、概率输出与评价归约连接到完整实验流程。最直接的认识是，标签的插值方式、ignore 屏蔽和共同评价网格会改变结果的含义，应先建立正确的计算口径，再讨论精度改善。分辨率提高与类别重加权对应不同的计算成本和类别取舍；多种子对照与实际效率测试帮助约束仅凭一次结果作出的判断。\n\n工程实现中的独立初始化、验证选模、恢复状态和来源哈希，使实验能够追溯到具体配置、数据与权重。项目达成的是一个完成复现、对照和离线使用的课程系统，其贡献在于实验组织和工程实现，架构与常规训练方法的原创性仍归原作者。\n"
    output += "\n## 参考文献\n\n" + get("参考资料")
    output += "\n\n## 附录 A 主要复现与演示入口\n\n先创建一个不存在的新训练目录，再取得同一固定版本数据和初始化编码器。不得在已有完整结果目录中直接覆盖重跑。以下入口构成复现顺序，具体安装和恢复条件见项目复现指南。\n\n```powershell\npython delivery/package.py --training-workspace ../segformer_retrain\ncd ../segformer_retrain\npython scripts/check_environment.py\npython scripts/download_camvid.py --workers 4\npython scripts/audit_dataset.py\npython scripts/check_model.py\npython scripts/smoke_test.py --offline --batch-size 4\npython scripts/check_overfit.py --offline --samples 20 --epochs 20\npython scripts/train.py --config configs/baseline.yaml --offline\npython scripts/run_experiments.py --offline\npython scripts/check_resolution.py --offline\npython scripts/train_resolution.py --offline\npython scripts/check_weighted.py --offline\npython scripts/train_weighted.py --offline\npython scripts/check_multiseed.py --offline\npython scripts/run_multiseed.py --offline\npython scripts/compare_multiseed.py\n```\n\nE4 正式运行前还需按 README 执行一次独立目录的真实 512×512 冒烟更新。恢复在相同训练命令后添加 --resume，要求配置、数据、代码和已保存状态一致。数据审计含旧绝对根路径时，必须在新训练目录重新验收。\n\n完整解压已有提交包后，可以直接验证并启动演示：\n\n```powershell\npython delivery/verify.py --output ../bundle_verification.json\npython demo/app.py --device auto --port 8765\n```\n\n浏览器访问 http://127.0.0.1:8765。Git 克隆只含代码与数值记录，最佳权重和照片示例在本地提交包中；依赖版本由 requirements.txt 固定。\n"
    output += "\n## 附录 B 结果与原始记录对应\n\n" + md_table(["正文证据", "原始记录"], [
        ["数据与真实诊断", "results/stage1/data_audit.json 与 overfit/overfit_check.json"],
        ["单种子结果", "results/E0_baseline 至 E5_weighted 的 metrics.json 与 train_log.csv"],
        ["分辨率与加权对照", "reports/resolution_ablation.md 与 weighted_ablation.md"],
        ["多种子统计与选择", "results/multiseed/runs.csv summary.csv selection.json"],
        ["效率测量", "results/benchmark_E4/benchmark.json"],
        ["图像来源与身份", "reports/report_evidence.json"],
        ["最终 ZIP 与离线验收", "delivery/release.json"],
        ["方法与来源许可", "THIRD_PARTY_NOTICES.md 与 docs/licenses/"]])
    output += "\n\n所有性能表来自同一批已完成实验。本报告保留单种子与种子均值、验证选模与测试评价、诊断与正式运行的区别。\n"
    output=output.replace("正式训练前的关键验收如下，诊断结果与正式 Val/Test 成绩严格区分。", "关键验收如下。前四项在正式训练前完成，最后一项为最终交付检查；诊断结果与正式 Val/Test 成绩严格区分。")
    output=output.replace("本次交付不自动追加这些实验，不把常规训练策略包装成新算法。", "这些方向尚未开展，本项目不据此报告额外实验成绩。")
    output=output.replace("报告数字由 `delivery/build_report.py` 从真实归档生成，图像原文件与复制件的哈希记入 reports/report_evidence.json。", "各实验数字对应完整归档，图像原文件与复制件的哈希记入 reports/report_evidence.json，综合报告与原始记录的对应关系记入 reports/word_report_evidence.json。")
    return output


def font(style, size, *, bold=False, east="宋体", latin="Times New Roman"):
    style.font.name=latin; style.font.size=Pt(size); style.font.bold=bold; style.font.italic=False; style.font.color.rgb=RGBColor(0,0,0)
    props=style.element.get_or_add_rPr()
    fonts=props.find(qn("w:rFonts"))
    if fonts is None: fonts=OxmlElement("w:rFonts"); props.insert(0,fonts)
    for key,value in (("ascii",latin),("hAnsi",latin),("eastAsia",east)):
        fonts.set(qn("w:"+key),value)
    for name in ("asciiTheme","hAnsiTheme","eastAsiaTheme","cstheme"):
        fonts.attrib.pop(qn("w:"+name),None)


def add_inline(paragraph, text):
    for item in re.split(r"(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)", text):
        if not item: continue
        link=re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)",item)
        if link:
            element=OxmlElement("w:hyperlink")
            element.set(qn("r:id"),paragraph.part.relate_to(link[2],"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True))
            run=OxmlElement("w:r"); props=OxmlElement("w:rPr")
            color=OxmlElement("w:color"); color.set(qn("w:val"),"000000"); props.append(color); run.append(props)
            node=OxmlElement("w:t"); node.text=link[1]; run.append(node); element.append(run); paragraph._p.append(element)
        else:
            run=paragraph.add_run(item.strip("*") if item.startswith("**") else item.strip("`") if item.startswith("`") else item)
            if item.startswith("**"): run.bold=True
            if item.startswith("`"): run.font.name="Consolas";run.font.size=Pt(9.5)


def math_text(text):
    run=OxmlElement("m:r"); t=OxmlElement("m:t");t.text=text;run.append(t);return run


def sub(base,index):
    element=OxmlElement("m:sSub");e=OxmlElement("m:e");e.append(math_text(base));s=OxmlElement("m:sub");s.append(math_text(index));element.extend([e,s]);return element


def fraction(numerator,denominator):
    element=OxmlElement("m:f");n=OxmlElement("m:num");d=OxmlElement("m:den")
    for child in numerator:n.append(child)
    for child in denominator:d.append(child)
    element.extend([n,d]);return element


def summation(index,expression):
    element=OxmlElement("m:nary");props=OxmlElement("m:naryPr")
    for name,value in (("chr","∑"),("limLoc","subSup"),("supHide","1")):
        node=OxmlElement("m:"+name);node.set(qn("m:val"),value);props.append(node)
    lower=OxmlElement("m:sub");lower.append(math_text(index));upper=OxmlElement("m:sup");e=OxmlElement("m:e");e.extend(expression)
    element.extend([props,lower,upper,e]);return element


def equation(doc,kind):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(10)
    math=OxmlElement("m:oMathPara");content=OxmlElement("m:oMath")
    if kind=="IOU":
        nodes=[sub("IoU","c"),math_text(" = "),fraction([sub("TP","c")],[sub("TP","c"),math_text(" + "),sub("FP","c"),math_text(" + "),sub("FN","c")])]
    elif kind=="CE":
        nodes=[sub("L","CE"),math_text(" = −"),fraction([summation("i∈V",[sub("w","yᵢ"),math_text(" ln "),sub("p","i,yᵢ")])],[summation("i∈V",[sub("w","yᵢ")])])]
    elif kind=="DICE":
        nodes=[sub("L","Dice"),math_text(" = 1 − "),fraction([math_text("1")],[math_text("|C|")]),summation("c∈C",[fraction([math_text("2"),summation("i∈V",[sub("p","ic"),sub("g","ic")]),math_text(" + ε")],[summation("i∈V",[sub("p","ic")]),math_text(" + "),summation("i∈V",[sub("g","ic")]),math_text(" + ε")])])]
    else:nodes=[math_text("L = "),sub("L","CE"),math_text(" + 0.5"),sub("L","Dice")]
    content.extend(nodes);math.append(content);p._p.append(math)


class Writer:
    def __init__(self):
        self.doc=Document();self.table_number=0;self.figure_number=0
        for style in self.doc.styles:
            for border in list(style.element.iter(qn("w:pBdr"))):border.getparent().remove(border)
        section=self.doc.sections[0]
        section.page_width=Inches(8.5);section.page_height=Inches(11)
        section.top_margin=Inches(.8);section.bottom_margin=Inches(.75)
        section.left_margin=Inches(.8);section.right_margin=Inches(.8)
        section.header_distance=Inches(.32);section.footer_distance=Inches(.32)
        section.different_first_page_header_footer=True
        for name,size,bold,east in (("Normal",11.5,False,"宋体"),("Title",22,True,"黑体"),("Subtitle",14,False,"宋体"),("Heading 1",15,True,"黑体"),("Heading 2",12.5,True,"黑体"),("Heading 3",11.5,True,"黑体"),("Caption",10,False,"宋体")):
            font(self.doc.styles[name],size,bold=bold,east=east)
        body=self.doc.styles["Normal"].paragraph_format
        body.line_spacing=1.35;body.space_after=Pt(6);body.first_line_indent=Pt(23);body.widow_control=True
        for name in ("Heading 1","Heading 2","Heading 3"):
            p=self.doc.styles[name].paragraph_format;p.first_line_indent=Pt(0);p.space_before=Pt(14);p.space_after=Pt(8);p.keep_with_next=True;p.keep_together=True
        cap=self.doc.styles["Caption"].paragraph_format;cap.first_line_indent=Pt(0);cap.space_before=Pt(4);cap.space_after=Pt(10);cap.keep_together=True
        for name in ("Code","TOC 1","TOC 2"):
            if name not in self.doc.styles:self.doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            font(self.doc.styles[name],10.5 if name!="Code" else 9.5,latin="Consolas" if name=="Code" else "Times New Roman")
            p=self.doc.styles[name].paragraph_format;p.first_line_indent=Pt(0);p.line_spacing=1.2;p.space_after=Pt(4)
            if name.startswith("TOC"):
                p.line_spacing=Pt(14);p.space_after=Pt(0);p.space_before=Pt(0)
                snap=OxmlElement("w:snapToGrid");snap.set(qn("w:val"),"0");self.doc.styles[name].element.get_or_add_pPr().append(snap)
        header=section.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        header.paragraph_format.first_line_indent=Pt(0)
        r=header.add_run("SegFormer B0 课程设计与实验综合报告");r.font.size=Pt(8.5)
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
        r=footer.add_run();begin=OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"),"begin");code=OxmlElement("w:instrText");code.text=" PAGE ";end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end");r._r.extend([begin,code,end])
        self.doc.core_properties.title="基于 SegFormer B0 的 CamVid 街景语义分割"
        self.doc.core_properties.subject="课程设计与实验综合报告"
        self.doc.core_properties.author="";self.doc.core_properties.last_modified_by=""
        self.doc.core_properties.comments=""

    def caption(self,text,table=False):
        if table:self.table_number+=1;label=f"表 {self.table_number}"
        else:self.figure_number+=1;label=f"图 {self.figure_number}"
        p=self.doc.add_paragraph(label+"  "+text,style="Caption");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if table:p.paragraph_format.keep_with_next=True
        return p

    def picture(self,path,description,width=6.4):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(0)
        shape=p.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set("descr",description)
        self.caption(description)

    def case(self,sample):
        for suffixes,labels in ((("input","gt"),("原图","真实标注 GT")),(("E0","E5"),("E0 预测","E5 预测"))):
            p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(0)
            for index,suffix in enumerate(suffixes):
                if index:p.add_run("  ")
                picture=p.add_run().add_picture(str(FIGURES / f"qualitative/{sample}_{suffix}.png"),width=Inches(3.28));picture._inline.docPr.set("descr",f"{sample} {labels[index]}")
            label=self.doc.add_paragraph(style="Caption");label.alignment=WD_ALIGN_PARAGRAPH.LEFT;label.paragraph_format.keep_with_next=True;label.paragraph_format.space_after=Pt(5)
            label.paragraph_format.tab_stops.add_tab_stop(Inches(1.64),WD_TAB_ALIGNMENT.CENTER)
            label.paragraph_format.tab_stops.add_tab_stop(Inches(5.01),WD_TAB_ALIGNMENT.CENTER)
            label.add_run("\t"+labels[0]+"\t"+labels[1])
        self.caption(f"固定测试样本 {sample} 的输入 标注与预测对照")

    def table(self,headers,rows):
        first=" ".join(headers)
        if "最佳轮次" in first and "实验" in first:title="单种子总体结果";widths=[1.25,.75,.6,.8,.8,.9,.8]
        elif headers[0]=="类别" and "E0" in headers:title="六组实验逐类测试 IoU";widths=[1.25]+[.92]*6
        elif "E5−E4 均值 pp" in first:title="三种子的逐类测试 IoU 均值与样本标准差";widths=[1.35,1.95,1.95,1.55]
        elif "训练网格" in first:title="冻结的单种子实验矩阵";widths=[.45,.85,.8,2.25,1.7]
        elif "原始 ignore" in first:title="数据划分及原始标签统计";widths=[1,1,1.4,2.2]
        elif headers[0]=="ID":title="类别顺序与实际训练权重";widths=[.45,1.7,1.7,1.65]
        elif "实际证据" in first:title="正式训练前与最终交付的验收记录";widths=[1.25,3,2.1]
        elif "种子数" in first:title="三种子总体均值与样本标准差";widths=[.45,.55,2,2,1.4]
        elif headers[0]=="种子":title="同种子下 E5 相对 E4 的配对差值";widths=[1.4,2.25,2.25]
        elif "方案" in headers and "最佳轮次" in headers:title="六次配对运行的完整结果";widths=[.55,.6,.75,1.4,1.4,1.3]
        elif "E4 推理精度" in first:
            title="E4 在本机的精度与推理效率";widths=[.65,.85,.95,.8,.75,1.1,1.25]
            headers=["精度","Val mIoU","模型均时\nms","P95\nms","模型\nFPS","图片流程\nms","显存峰值\nMiB"]
        else:title="正文结果与原始归档对应";widths=[1.7,4.7] if len(headers)==2 else [6.6/len(headers)]*len(headers)
        if len(widths)!=len(headers):widths=[6.6/len(headers)]*len(headers)
        scale=min(1,6.8/sum(widths));widths=[value*scale for value in widths]
        self.caption(title,table=True)
        table=self.doc.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        for column,width in zip(table.columns,widths):column.width=Inches(width)
        for row_index,(row_values,row) in enumerate([(headers,table.rows[0]),*[(values,table.add_row()) for values in rows]]):
            is_header=row._tr is table._tbl.tr_lst[0]
            props=row._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");props.append(cant)
            if is_header:repeat=OxmlElement("w:tblHeader");props.append(repeat)
            for index,(cell,text) in enumerate(zip(row.cells,row_values)):
                cell.width=Inches(widths[index]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cellprops=cell._tc.get_or_add_tcPr();margins=OxmlElement("w:tcMar")
                for side in ("top","left","bottom","right"):
                    element=OxmlElement("w:"+side);element.set(qn("w:w"),"80");element.set(qn("w:type"),"dxa");margins.append(element)
                cellprops.append(margins)
                if is_header:fill=OxmlElement("w:shd");fill.set(qn("w:fill"),"E4EBF3");cellprops.append(fill)
                p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.keep_with_next=(row_index==0 or (len(rows)<=7 and row_index<len(rows)))
                narrative=len(headers)==2 or "实际证据" in headers
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if (narrative or (index==0 and headers[0] not in ("ID","种子","方案","组别"))) else WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p,str(text))
                for run in p.runs:run.font.size=Pt(10.5);run.bold=is_header
        borders=OxmlElement("w:tblBorders")
        for side in ("top","left","bottom","right","insideH","insideV"):
            element=OxmlElement("w:"+side);element.set(qn("w:val"),"single");element.set(qn("w:sz"),"4");element.set(qn("w:color"),"D9D9D9");borders.append(element)
        table._tbl.tblPr.append(borders)
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.add_run().font.size=Pt(3)

    def toc(self):
        self.doc.add_page_break();p=self.doc.add_paragraph("目录",style="TOC Heading");font(self.doc.styles["TOC Heading"],16,bold=True,east="黑体")
        p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
        begin=OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"),"begin")
        code=OxmlElement("w:instrText");code.set(qn("xml:space"),"preserve");code.text=' TOC \\o "1-2" \\h \\z \\u '
        separate=OxmlElement("w:fldChar");separate.set(qn("w:fldCharType"),"separate");end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end")
        p.add_run()._r.extend([begin,code,separate,end]);self.doc.add_page_break()

    def write(self,markdown):
        lines=markdown.splitlines();index=0;cover_lines=0
        while index<len(lines):
            line=lines[index].strip()
            if not line:index+=1;continue
            if line.startswith("# "):
                p=self.doc.add_paragraph(line[2:],style="Title");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(30);p.paragraph_format.space_after=Pt(20)
            elif cover_lines<3 and not line.startswith("##"):
                p=self.doc.add_paragraph(line,style="Subtitle" if cover_lines==0 else "Normal");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(10);cover_lines+=1
            elif line=="[[TOC]]":self.toc()
            elif line.startswith("[[EQUATION:"):equation(self.doc,line[11:-2])
            elif line.startswith("[[CASE:"):self.case(line[7:-2])
            elif line.startswith("##"):
                level=len(line)-len(line.lstrip("#"))-1;text=line.lstrip("# ").strip()
                p=self.doc.add_paragraph(text,style="Heading "+str(level))
                if text.startswith(("7 定性","8 推理","附录 A","附录 B")) or text in ("街景样本二","参考文献"):p.paragraph_format.page_break_before=True
            elif line.startswith("```"):
                index+=1
                while index<len(lines) and not lines[index].startswith("```"):
                    p=self.doc.add_paragraph(lines[index],style="Code");p.paragraph_format.keep_with_next=index+1<len(lines) and not lines[index+1].startswith("```");index+=1
            elif line.startswith("|"):
                block=[]
                while index<len(lines) and lines[index].startswith("|"):
                    cells=[cell.strip() for cell in lines[index].strip("|").split("|")]
                    if not all(re.fullmatch(r":?-+:?",cell) for cell in cells):block.append(cells)
                    index+=1
                self.table(block[0],block[1:]);continue
            elif re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)",line):
                image=re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)",line)
                self.picture(ROOT/"reports"/image[2],image[1],5.5 if "混淆" in image[1] else 6.4)
            else:
                p=self.doc.add_paragraph(style="Normal")
                if line.startswith("- "):
                    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Pt(15)
                    # Numbered bibliography with full source details, no stage-progress fragments.
                    line=line[2:]
                    if any(name in line for name in ("Xie E.","Brostow G.","Hugging Face.","NVIDIA /", "lih627.")):
                        number=getattr(self,"reference_number",0)+1;self.reference_number=number;line=f"[{number}] "+line
                add_inline(p,line)
                following=next((value.strip() for value in lines[index+1:] if value.strip()),"")
                if following.startswith("|"):p.paragraph_format.keep_with_next=True
            index+=1
        self.doc.save(OUTPUT)


def main():
    records=verified_inputs();markdown=merged_markdown(records)
    working=ROOT/".cache/word_report";working.mkdir(parents=True,exist_ok=True)
    (working/"combined_source.md").write_text(markdown,encoding="utf-8")
    writer=Writer();writer.write(markdown)
    evidence={"output":OUTPUT.relative_to(ROOT).as_posix(),"verified_runs":len(records),"tables":writer.table_number,"figures":writer.figure_number,
              "source_documents":{name:digest(ROOT/name) for name in SOURCE_DOCS},
              "metrics":{f"results/{record['name']}/metrics.json":digest(ROOT/"results"/record["name"]/"metrics.json") for record in records},
              "test_inference_executed":False,"full_training_executed":False}
    (ROOT/"reports/word_report_evidence.json").write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"output":str(OUTPUT),"characters":len(markdown),"tables":writer.table_number,"figures":writer.figure_number},ensure_ascii=False,indent=2))


if __name__=="__main__":main()

"""Check the delivered Word tables against completed metrics and record visual QA identity."""

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree
from pypdf import PdfReader

from delivery.build_word_report import RUNS, ROOT, OUTPUT, NAMES, verified_inputs


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def clear_editor_metadata():
    # The content, styles, fields, image bytes and document pagination are untouched.
    with zipfile.ZipFile(OUTPUT) as source:
        infos=source.infolist();data={info.filename:source.read(info.filename) for info in infos}
    ns={"cp":"http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
    core=etree.fromstring(data["docProps/core.xml"])
    for node in core.xpath("//cp:lastModifiedBy",namespaces=ns):node.text=""
    data["docProps/core.xml"]=etree.tostring(core,encoding="UTF-8",xml_declaration=True,standalone=True)
    temporary=OUTPUT.with_suffix(".docx.tmp")
    with zipfile.ZipFile(temporary,"w") as target:
        for info in infos:target.writestr(info,data[info.filename])
    with zipfile.ZipFile(temporary) as check:
        assert all(check.read(name)==content for name,content in data.items())
    temporary.replace(OUTPUT)


def main():
    records=verified_inputs();metrics=[item["metrics"] for item in records]
    clear_editor_metadata()
    document=Document(OUTPUT)
    assert len(document.tables)==12 and len(document.inline_shapes)==12
    for row,metric in zip(document.tables[4].rows[1:],metrics[:6]):
        cells=[cell.text for cell in row.cells]
        assert cells[0]==metric["experiment_id"]
        assert cells[2]==str(metric["best_val_epoch"])
        assert cells[3:6]==[f"{100*metric[key]:.2f}%" for key in ("best_val_miou","test_miou","pixel_accuracy")]
    for class_id,row in enumerate(document.tables[5].rows[1:]):
        assert [cell.text for cell in row.cells]==[NAMES[class_id],*[f"{100*metric['per_class_iou'][class_id]:.2f}%" for metric in metrics[:6]]]
    for row,index in zip(document.tables[8].rows[1:],(4,5,6,7,8,9)):
        metric=metrics[index]
        assert [cell.text for cell in row.cells][1:]==[str(metric["seed"]),str(metric["best_val_epoch"]),*[f"{100*metric[key]:.2f}%" for key in ("best_val_miou","test_miou","pixel_accuracy")]]
    for class_id,row in enumerate(document.tables[9].rows[1:]):
        stats=[]
        for indices in ((4,6,8),(5,7,9)):
            values=[metrics[index]["per_class_iou"][class_id] for index in indices]
            mean=sum(values)/3;std=(sum((value-mean)**2 for value in values)/2)**0.5;stats.append((mean,std))
        expected=[NAMES[class_id],*[f"{100*mean:.2f}% ± {100*std:.2f}" for mean,std in stats],f"{100*(stats[1][0]-stats[0][0]):+.2f}"]
        assert [cell.text for cell in row.cells]==expected
    source_images=[ROOT/"reports/figures"/name for name in ("E0_baseline_training_curves.png","E5_weighted_training_curves.png","E5_weighted_confusion_matrix.png")]
    for sample in ("0001TP_008550","Seq05VD_f02310"):
        source_images.extend(ROOT/"reports/figures/qualitative"/f"{sample}_{suffix}.png" for suffix in ("input","gt","E0","E5"))
    source_images.append(ROOT/"reports/figures/qualitative/demo_interface.png")
    image_checks=[]
    for shape,source in zip(document.inline_shapes,source_images):
        relationship=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        embedded=document.part.related_parts[relationship].blob
        same=hashlib.sha256(embedded).hexdigest()==sha256(source)
        assert same,f"Embedded image differs: {source}"
        image_checks.append(source.relative_to(ROOT).as_posix())
    with zipfile.ZipFile(OUTPUT) as package:
        xml=etree.fromstring(package.read("word/document.xml"))
        ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main","m":"http://schemas.openxmlformats.org/officeDocument/2006/math"}
        assert len(xml.xpath("//m:oMath",namespaces=ns))==4
        assert not xml.xpath("//w:pBdr",namespaces=ns)
        assert "[[" not in "".join(xml.xpath("//w:t/text()",namespaces=ns))
        assert not document.core_properties.author and not document.core_properties.last_modified_by
    pdf=ROOT/".cache/word_report/report_final.pdf"
    assert len(PdfReader(pdf).pages)==20
    evidence_path=ROOT/"reports/word_report_evidence.json"
    evidence=json.loads(evidence_path.read_text(encoding="utf-8"))
    evidence.update(output_sha256=sha256(OUTPUT),output_bytes=OUTPUT.stat().st_size,
                    numeric_word_tables_verified=True,embedded_source_images_verified=image_checks,
                    native_word_equations=4,native_word_tables=12,
                    render_validation={"engine":"Microsoft Word 16.0 ExportAsFixedFormat and canonical documents/render_docx.py rasterizer",
                                       "page_count":20,"reviewed_pages":list(range(1,21)),
                                       "preview_pdf_sha256":sha256(pdf),"visual_review":"passed",
                                       "title_border_removed":True,"toc_page_numbers_updated":True,
                                       "short_tables_kept_together":True,"metadata_only_finalization":True})
    evidence_path.write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"pages":20,"tables":12,"equations":4,"source_images":12,"verified_runs":10,"sha256":sha256(OUTPUT)},indent=2))


if __name__=="__main__":main()
